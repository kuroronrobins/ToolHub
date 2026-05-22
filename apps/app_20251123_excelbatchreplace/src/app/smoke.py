from __future__ import annotations

import importlib
from typing import Iterable

from app.paths import setting_workbook_path, user_cache_dir, user_data_root, user_templates_dir

APP_MODULES: tuple[str, ...] = (
    "app.config",
    "app.excel_controller",
    "app.batch_logic",
    "app.settings_store",
    "app.ui_main",
)

RUNTIME_MODULES: tuple[str, ...] = (
    "win32com.client",
    "openpyxl",
)


def _import_modules(names: Iterable[str], errors: list[str]) -> None:
    for name in names:
        try:
            importlib.import_module(name)
        except Exception as exc:  # noqa: BLE001
            errors.append(f"import failed: {name}: {exc}")


def _check_setting_workbook(errors: list[str]) -> None:
    path = setting_workbook_path()
    if not path.exists():
        errors.append(f"missing required asset: {path}")
        return
    try:
        from openpyxl import load_workbook  # type: ignore

        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            if "AdjustLine" not in wb.sheetnames:
                errors.append(f"Setting.xlsx missing sheet: AdjustLine ({path})")
        finally:
            wb.close()
    except Exception as exc:  # noqa: BLE001
        errors.append(f"Setting.xlsx read failed: {path}: {exc}")


def run_smoke_check() -> int:
    errors: list[str] = []
    _import_modules(RUNTIME_MODULES, errors)
    _import_modules(APP_MODULES, errors)
    _check_setting_workbook(errors)

    try:
        user_data_root()
        user_cache_dir()
        user_templates_dir()
    except Exception as exc:  # noqa: BLE001
        errors.append(f"user data directory check failed: {exc}")

    if errors:
        print("[toolhub-smoke] failed")
        for error in errors:
            print(f"- {error}")
        return 1

    print("[toolhub-smoke] ok")
    return 0
